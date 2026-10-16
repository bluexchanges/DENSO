from pathlib import Path
import json,re,hashlib,collections
from pptx import Presentation

ROOT=Path('F:/Denso')
OUT=ROOT/'docs_denso_hack'/'_audit'
OUT.mkdir(parents=True,exist_ok=True)
deck=ROOT/'Denso-Hack_Presentation-EDABK_TriH.pptx'
p=Presentation(deck)
slides=[]
def shape_text(s):
    out=[]
    if s.has_text_frame: out.append(s.text)
    if s.has_table:
        out.extend(' | '.join(c.text for c in row.cells) for row in s.table.rows)
    if s.shape_type==6:
        for ch in s.shapes: out.extend(shape_text(ch))
    if getattr(s,'has_chart',False):
        try: out.append('CHART: '+str(s.chart.chart_type))
        except Exception: pass
    return out
for i,s in enumerate(p.slides,1):
    texts=[]
    for sh in s.shapes: texts.extend(shape_text(sh))
    notes=s.notes_slide.notes_text_frame.text if s.has_notes_slide else ''
    slides.append({'slide':i,'text':texts,'notes':notes,'pictures':sum(1 for x in s.shapes if x.shape_type==13)})
(OUT/'slides_extracted.json').write_text(json.dumps(slides,ensure_ascii=False,indent=2),encoding='utf-8')
(OUT/'slides_extracted.txt').write_text('\n\n'.join(f"SLIDE {s['slide']}\n"+'\n'.join(s['text'])+'\nNOTES: '+s['notes'] for s in slides),encoding='utf-8')
inventory=[{'path':str(f),'bytes':f.stat().st_size,'sha256':hashlib.sha256(f.read_bytes()).hexdigest()} for f in ROOT.iterdir() if f.is_file()]
(OUT/'initial_inventory.json').write_text(json.dumps(inventory,ensure_ascii=False,indent=2),encoding='utf-8')
codex=Path('C:/Users/ADMIN/.codex')
pat=re.compile(r'RH20T|MuJoCo|\bPanda\b|\bDenso\b|EDABK|\bCVAE\b|temporal.ensembl|robot.learning|overfit',re.I)
files=sorted(list((codex/'sessions').rglob('*.jsonl'))+list((codex/'archived_sessions').rglob('*.jsonl')))
hits=[];summary=[]
current='01a145e5-4162-7782-a6a6-2403a7a23af8'
for f in files:
    if current in f.name: continue
    cwd='';counts=collections.Counter(); entries=[]
    for n,line in enumerate(f.open(encoding='utf-8'),1):
        try: obj=json.loads(line)
        except Exception: continue
        pl=obj.get('payload',{})
        if obj.get('type')=='session_meta': cwd=pl.get('cwd','')
        kind=obj.get('type'); role=pl.get('role','')
        if kind=='response_item' and pl.get('type')=='message' and role in ('user','assistant'):
            txt='\n'.join(c.get('text','') for c in pl.get('content',[]) if c.get('type') in ('input_text','output_text','text'))
        elif kind=='response_item' and pl.get('type') in ('function_call_output','custom_tool_call_output'):
            txt=str(pl.get('output',''));role='tool'
        else: continue
        # Ignore injected environment and skill catalog matches.
        if role=='user' and ('<environment_context>' in txt or '<skills_instructions>' in txt or '# AGENTS.md instructions' in txt): continue
        if pat.search(txt):
            counts[role]+=1
            entries.append({'file':str(f),'line':n,'time':obj.get('timestamp'),'cwd':cwd,'role':role,'text':txt})
    if entries:
        hits.extend(entries);summary.append({'file':str(f),'cwd':cwd,'hits':dict(counts)})
(OUT/'history_hits.json').write_text(json.dumps(hits,ensure_ascii=False,indent=2),encoding='utf-8')
(OUT/'history_index.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
print('SLIDES',len(slides),'NOTES_NONEMPTY',sum(bool(s['notes'].strip()) for s in slides))
print('HISTORY_FILES_SCANNED',len(files)-1,'RELEVANT_CANDIDATES',len(summary))
for s in summary: print(json.dumps(s,ensure_ascii=False))
